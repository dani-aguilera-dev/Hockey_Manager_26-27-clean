import json
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TOOLS = Path(__file__).resolve().parent
BASE = TOOLS.parent
SRC = str(BASE / "equipos_hockey_hierba_espana_masculino_2026_27_con_posiciones.xlsx")
OUT = str(BASE / "equipos_hockey_hierba_espana_masculino_2026_27_con_atributos.xlsx")
CSV_OUT = str(BASE / "jugadores_atributos_2026_27.csv")
df = pd.read_csv(TOOLS / "player_attributes_v2.csv", encoding="utf-8-sig").fillna("")
df["edad"] = pd.to_numeric(df["edad"], errors="coerce").fillna(0).astype(int)
df["goles_num"] = pd.to_numeric(df["goles"], errors="coerce").fillna(0).astype(int)
for col in ["pos_url", "pos_fecha", "pos_nota"]:
    if col not in df.columns:
        df[col] = ""
    df[col] = df[col].fillna("")

ATTRS = [
    "control",
    "pase",
    "conduccion",
    "golpeo",
    "flick",
    "regate_3d",
    "definicion",
    "aerial",
    "pc_injection",
    "pc_trap",
    "pc_flick",
    "pc_deflection",
    "pc_first_runner",
    "pc_gk_save",
    "stroke_shot",
    "stroke_save",
    "free_hit",
    "posicionamiento",
    "pressing",
    "cobertura",
    "transicion",
    "desmarque",
    "velocidad",
    "resistencia",
    "fuerza",
    "agilidad",
    "coordinacion",
    "decisiones",
    "concentracion",
    "competitividad",
    "liderazgo",
    "reflejos",
    "blocaje",
    "distribucion",
    "uno_vs_uno",
    "alto",
    "rebote_gk",
]
GROUP_OF = {}
for g, arr in {
    "Técnica": ["control", "pase", "conduccion", "golpeo", "flick", "regate_3d", "definicion"],
    "3D y set pieces": [
        "aerial",
        "pc_injection",
        "pc_trap",
        "pc_flick",
        "pc_deflection",
        "pc_first_runner",
        "pc_gk_save",
        "stroke_shot",
        "stroke_save",
        "free_hit",
    ],
    "Táctica": ["posicionamiento", "pressing", "cobertura", "transicion", "desmarque"],
    "Física": ["velocidad", "resistencia", "fuerza", "agilidad", "coordinacion"],
    "Mental": ["decisiones", "concentracion", "competitividad", "liderazgo"],
    "Portero": ["reflejos", "blocaje", "distribucion", "uno_vs_uno", "alto", "rebote_gk"],
}.items():
    for a in arr:
        GROUP_OF[a] = g
LABEL = {
    "control": "Control",
    "pase": "Pase",
    "conduccion": "Conducción",
    "golpeo": "Golpeo (hit/slap/sweep)",
    "flick": "Flick / scoop",
    "regate_3d": "Regate 3D",
    "definicion": "Definición",
    "aerial": "Aerial / scoop",
    "pc_injection": "Saque de PC (inyección)",
    "pc_trap": "Parada de PC (stopper)",
    "pc_flick": "Tiro de PC / drag flick",
    "pc_deflection": "Desvío en PC",
    "pc_first_runner": "Primer runner / defensa PC",
    "pc_gk_save": "Parada de PC (portero)",
    "stroke_shot": "Tiro de stroke",
    "stroke_save": "Parada de stroke",
    "free_hit": "Faltas, 16 yardas y córners",
    "posicionamiento": "Posicionamiento",
    "pressing": "Pressing",
    "cobertura": "Cobertura",
    "transicion": "Transición",
    "desmarque": "Desmarque",
    "velocidad": "Velocidad",
    "resistencia": "Resistencia",
    "fuerza": "Fuerza",
    "agilidad": "Agilidad",
    "coordinacion": "Coordinación",
    "decisiones": "Decisiones",
    "concentracion": "Concentración",
    "competitividad": "Competitividad",
    "liderazgo": "Liderazgo",
    "reflejos": "Reflejos",
    "blocaje": "Blocaje",
    "distribucion": "Distribución",
    "uno_vs_uno": "1 contra 1",
    "alto": "Paradas altas",
    "rebote_gk": "Control de rebotes",
}
DESC = {
    "control": "Control y recepción en estático y en movimiento.",
    "pase": "Pase de push, hit y slap; precisión y peso.",
    "conduccion": "Conducción con la bola pegada al stick.",
    "golpeo": "Golpeo (hit/slap/sweep) y potencia.",
    "flick": "Flick y scoop para elevar la bola.",
    "regate_3d": "Regate 3D, cambios de plano y espacios reducidos.",
    "definicion": "Definición en área, uno contra uno y remate.",
    "aerial": "Aerial/scoop: envío y recepción de bolas elevadas.",
    "pc_injection": "Saque de penalty corner: precisión, velocidad y consistencia.",
    "pc_trap": "Parada/control del saque de PC (stopper).",
    "pc_flick": "Tiro de penalty corner: drag flick, flick directo y potencia.",
    "pc_deflection": "Desvío en PC: timing, mano y colocación.",
    "pc_first_runner": "Primer runner: salida defensiva de PC y presión.",
    "pc_gk_save": "Parada de PC del portero.",
    "stroke_shot": "Tiro de penalty stroke.",
    "stroke_save": "Parada de penalty stroke.",
    "free_hit": "Faltas, 16 yardas, córners y saques de banda.",
    "posicionamiento": "Lectura del juego y ocupación de espacios.",
    "pressing": "Presión, recuperación y pressing tras pérdida.",
    "cobertura": "Cobertura defensiva, ayudas y basculación.",
    "transicion": "Transición ataque-defensa y defensa-ataque.",
    "desmarque": "Desmarque, ruptura y generación de espacios.",
    "velocidad": "Velocidad de reacción y desplazamiento.",
    "resistencia": "Capacidad aeróbica y repetir esfuerzos.",
    "fuerza": "Fuerza aplicada a golpeo, duelos y PC.",
    "agilidad": "Agilidad, cambios de dirección y coordinación.",
    "coordinacion": "Coordinación óculo-manual y técnica corporal.",
    "decisiones": "Toma de decisiones bajo presión.",
    "concentracion": "Concentración y consistencia.",
    "competitividad": "Intensidad, carácter y competitividad.",
    "liderazgo": "Liderazgo y comunicación.",
    "reflejos": "Reflejos y reacción (portero).",
    "blocaje": "Blocaje y control de paradas (portero).",
    "distribucion": "Distribución y saque del portero.",
    "uno_vs_uno": "Salidas y 1 contra 1 (portero).",
    "alto": "Paradas altas y juego aéreo (portero).",
    "rebote_gk": "Control de rebotes y segunda jugada (portero).",
}
SPEC_LABEL = {
    "pc_flicker": "Drag flicker / tiro de PC",
    "pc_injector": "Inyector de PC",
    "pc_stopper": "Stopper de PC",
    "pc_deflector": "Desviador de PC",
    "pc_first_runner": "Primer runner",
    "pc_goalkeeper": "Portero de PC",
    "stroke_taker": "Tirador de stroke",
    "stroke_goalkeeper": "Portero de stroke",
    "aerial_specialist": "Especialista aéreo",
    "free_hit_taker": "Especialista en faltas",
}

# team specialists map
team_spec = {}
for (liga, equipo), g in df.groupby(["liga", "equipo"]):
    j = json.loads(str(g.iloc[0].get("team_specialists", "") or "{}"))
    team_spec[(liga, equipo)] = j
# invert player -> roles
player_roles = defaultdict(list)
for k, spec in team_spec.items():
    for role, name in spec.items():
        player_roles[(k[0], k[1], name)].append(role)

# CSV full
out = df.copy()
out["subposicion"] = out["subposicion"]
out["grupo"] = out["position_group"]
out["arquetipo"] = out["archetype"]
out["overall_principal"] = out["overall"]
out["potencial"] = out["potential"]
out["valor_eur"] = out["value"]
out["polivalencia"] = out["polivalencia"]
for a in ATTRS:
    out[LABEL[a]] = out[f"attr_{a}"]
out["especialista_s"] = out.apply(
    lambda r: "; ".join(SPEC_LABEL.get(x, x) for x in player_roles.get((r["liga"], r["equipo"], r["jugador"]), [])),
    axis=1,
)
csv_cols = (
    [
        "liga",
        "equipo",
        "club",
        "jugador",
        "perfil",
        "edad",
        "goles",
        "pos",
        "subposicion",
        "position_group",
        "src",
        "conf",
        "archetype",
        "overall",
        "potential",
        "value",
        "polivalencia",
        "overall_portero",
        "overall_central",
        "overall_lateral",
        "overall_medio",
        "overall_delantero",
    ]
    + [f"attr_{a}" for a in ATTRS]
    + ["especialista_s"]
)
out[csv_cols].to_csv(CSV_OUT, index=False, encoding="utf-8-sig")

# Load existing workbook to preserve images/formatting
wb = openpyxl.load_workbook(SRC)
header_fill = PatternFill("solid", fgColor="123B5D")
header_font = Font(color="FFFFFF", bold=True, size=10)
sub_fill = PatternFill("solid", fgColor="1D5C7A")
thin = Side(style="thin", color="D0D7DE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

new_cols = [
    ("Subposición", 17),
    ("Grupo", 18),
    ("Arquetipo", 19),
    ("Overall", 20),
    ("Potencial", 21),
    ("Valor (€)", 22),
    ("Polivalencia", 23),
    ("Overall POR", 24),
    ("Overall CEN", 25),
    ("Overall LAT", 26),
    ("Overall MED", 27),
    ("Overall DEL", 28),
    ("Especialista(s)", 29),
    ("URL posición", 30),
    ("Fecha posición", 31),
]
# add headers to three player sheets
for sh in ["Jugadores DHA", "Jugadores DHB", "San Fernando 1ª Andaluza"]:
    ws = wb[sh]
    for name, col in new_cols:
        c = ws.cell(row=1, column=col, value=name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border
        ws.column_dimensions[get_column_letter(col)].width = {
            "Subposición": 13,
            "Grupo": 12,
            "Arquetipo": 22,
            "Overall": 9,
            "Potencial": 9,
            "Valor (€)": 12,
            "Polivalencia": 12,
            "Overall POR": 10,
            "Overall CEN": 10,
            "Overall LAT": 10,
            "Overall MED": 10,
            "Overall DEL": 10,
            "Especialista(s)": 34,
            "URL posición": 45,
            "Fecha posición": 14,
        }.get(name, 12)
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(31)}{ws.max_row}"
    # map rows
    rows = {}
    for r in range(2, ws.max_row + 1):
        key = (
            str(ws.cell(r, 2).value or "").strip(),
            str(ws.cell(r, 3).value or "").strip(),
            str(ws.cell(r, 5).value or "").strip(),
            str(ws.cell(r, 6).value or "").strip(),
        )
        rows[key] = r
    for _, d in df.iterrows():
        key = (str(d["liga"]).strip(), str(d["equipo"]).strip(), str(d["jugador"]).strip(), str(d["perfil"]).strip())
        r = rows.get(key)
        if not r:
            # fallback by liga/equipo/jugador
            for k2, r2 in rows.items():
                if k2[0] == key[0] and k2[1] == key[1] and k2[2] == key[2]:
                    r = r2
                    break
        if not r:
            continue
        roles = player_roles.get((d["liga"], d["equipo"], d["jugador"]), [])
        vals = [
            d["subposicion"],
            d["position_group"],
            d["archetype"],
            float(d["overall"]),
            float(d["potential"]),
            int(float(d["value"])),
            float(d["polivalencia"]),
            float(d["overall_portero"]),
            float(d["overall_central"]),
            float(d["overall_lateral"]),
            float(d["overall_medio"]),
            float(d["overall_delantero"]),
            "; ".join(SPEC_LABEL.get(x, x) for x in roles),
            d.get("pos_url", ""),
            d.get("pos_fecha", ""),
        ]
        for (name, col), v in zip(new_cols, vals):
            c = ws.cell(row=r, column=col, value=v)
            c.border = border
            if name in (
                "Overall",
                "Potencial",
                "Overall POR",
                "Overall CEN",
                "Overall LAT",
                "Overall MED",
                "Overall DEL",
                "Polivalencia",
            ):
                c.number_format = "0.0"
                c.alignment = center
            elif name == "Valor (€)":
                c.number_format = "#,##0"
                c.alignment = center
            elif name == "Especialista(s)":
                c.alignment = left
            else:
                c.alignment = center
    # color scale overall
    from openpyxl.formatting.rule import ColorScaleRule

    ws.conditional_formatting.add(
        f"T2:T{ws.max_row}",
        ColorScaleRule(
            start_type="num",
            start_value=55,
            start_color="F8696B",
            mid_type="num",
            mid_value=72,
            mid_color="FFEB84",
            end_type="num",
            end_value=90,
            end_color="63BE7B",
        ),
    )

# New sheet: Atributos
ws = wb.create_sheet("Atributos")
headers = [
    "Liga",
    "Equipo",
    "Jugador",
    "Perfil",
    "Edad",
    "Goles",
    "Subposición",
    "Arquetipo",
    "Overall",
    "Potencial",
    "Valor (€)",
] + [LABEL[a] for a in ATTRS]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"
for i, d in df.iterrows():
    vals = [
        d["liga"],
        d["equipo"],
        d["jugador"],
        d["perfil"],
        int(d["edad"]),
        d["goles"],
        d["subposicion"],
        d["archetype"],
        float(d["overall"]),
        float(d["potential"]),
        int(float(d["value"])),
    ] + [int(d[f"attr_{a}"]) for a in ATTRS]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i + 2, column=j, value=v)
        c.border = border
        if j >= 5:
            c.alignment = center
        if j in (9, 10):
            c.number_format = "0.0"
        if j == 11:
            c.number_format = "#,##0"
        if j == 3:
            c.alignment = left
for j in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 13 if j > 4 else (22 if j == 3 else 11)
    if j >= 12:
        ws.column_dimensions[get_column_letter(j)].width = 12
# color scale attributes
from openpyxl.formatting.rule import ColorScaleRule

for j in range(12, len(headers) + 1):
    ws.conditional_formatting.add(
        f"{get_column_letter(j)}2:{get_column_letter(j)}{len(df) + 1}",
        ColorScaleRule(
            start_type="num",
            start_value=30,
            start_color="F8696B",
            mid_type="num",
            mid_value=65,
            mid_color="FFEB84",
            end_type="num",
            end_value=95,
            end_color="63BE7B",
        ),
    )

# Sheet Especialistas
ws = wb.create_sheet("Especialistas")
roles = list(SPEC_LABEL.keys())
headers = ["Liga", "Equipo"] + [SPEC_LABEL[r] for r in roles]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
teams = []
for (liga, equipo), g in df.groupby(["liga", "equipo"]):
    teams.append((liga, equipo, g.iloc[0]["club"]))
for i, (liga, equipo, club) in enumerate(teams, 2):
    spec = team_spec.get((liga, equipo), {})
    vals = [liga, equipo] + [spec.get(r, "") for r in roles]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = border
        c.alignment = left if j <= 2 else center
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(teams) + 1}"
for j in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 26 if j == 2 else 20

# Sheet Equipos atributos
ws = wb.create_sheet("Equipos atributos")
headers = [
    "Liga",
    "Equipo",
    "Club",
    "Jugadores",
    "Invitados",
    "Edad media",
    "Goles",
    "Overall medio",
    "Potencial medio",
    "Valor total (€)",
    "Técnica",
    "3D y set pieces",
    "Táctica",
    "Física",
    "Mental",
    "Portero",
    "Mejor jugador",
]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
for i, (liga, equipo, club) in enumerate(teams, 2):
    g = df[(df.liga == liga) & (df.equipo == equipo)]
    sq = g[g.perfil == "Jugador"]
    base = sq if len(sq) else g
    cat = {}
    for catname, arr in {
        "Técnica": ["control", "pase", "conduccion", "golpeo", "flick", "regate_3d", "definicion"],
        "3D y set pieces": [
            "aerial",
            "pc_injection",
            "pc_trap",
            "pc_flick",
            "pc_deflection",
            "pc_first_runner",
            "pc_gk_save",
            "stroke_shot",
            "stroke_save",
            "free_hit",
        ],
        "Táctica": ["posicionamiento", "pressing", "cobertura", "transicion", "desmarque"],
        "Física": ["velocidad", "resistencia", "fuerza", "agilidad", "coordinacion"],
        "Mental": ["decisiones", "concentracion", "competitividad", "liderazgo"],
        "Portero": ["reflejos", "blocaje", "distribucion", "uno_vs_uno", "alto", "rebote_gk"],
    }.items():
        cat[catname] = round(sum(g[f"attr_{a}"].mean() for a in arr) / len(arr), 1)
    best = g.nlargest(1, "overall").iloc[0]
    vals = [
        liga,
        equipo,
        club,
        len(sq),
        len(g) - len(sq),
        round(base["edad"].mean(), 1) if len(base) else 0,
        int(sq["goles_num"].sum()),
        round(base["overall"].mean(), 1),
        round(base["potential"].mean(), 1),
        int(g["value"].sum()),
        cat["Técnica"],
        cat["3D y set pieces"],
        cat["Táctica"],
        cat["Física"],
        cat["Mental"],
        cat["Portero"],
        f"{best['jugador']} ({best['overall']})",
    ]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = border
        c.alignment = left if j in (2, 3, 17) else center
        if j in (8, 9):
            c.number_format = "0.0"
        if j == 10:
            c.number_format = "#,##0"
        if j >= 11 and j <= 16:
            c.number_format = "0.0"
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:Q{len(teams) + 1}"
for j in range(1, 18):
    ws.column_dimensions[get_column_letter(j)].width = 18 if j not in (3, 17) else 30
# color scale team overall
ws.conditional_formatting.add(
    f"H2:H{len(teams) + 1}",
    ColorScaleRule(
        start_type="num",
        start_value=55,
        start_color="F8696B",
        mid_type="num",
        mid_value=68,
        mid_color="FFEB84",
        end_type="num",
        end_value=85,
        end_color="63BE7B",
    ),
)

# Sheet Diccionario
ws = wb.create_sheet("Diccionario atributos")
headers = ["Grupo", "Atributo", "Etiqueta", "Descripción"]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
for i, a in enumerate(ATTRS, 2):
    vals = [GROUP_OF[a], a, LABEL[a], DESC[a]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = border
        c.alignment = left if j == 4 else center
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:D{len(ATTRS) + 1}"
for j, w in enumerate([16, 18, 30, 65], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# Sheet Cobertura atributos
if "Cobertura atributos" in wb.sheetnames:
    del wb["Cobertura atributos"]
ws = wb.create_sheet("Cobertura atributos")
r = 1
ws.cell(r, 1, "Cobertura del modelo de atributos").font = Font(bold=True, size=14, color="123B5D")
r += 2
ws.cell(r, 1, "Perfiles totales")
ws.cell(r, 2, len(df))
r += 1
ws.cell(r, 1, "Jugadores")
ws.cell(r, 2, int((df.perfil == "Jugador").sum()))
r += 1
ws.cell(r, 1, "Invitados")
ws.cell(r, 2, int((df.perfil == "Invitado").sum()))
r += 1
ws.cell(r, 1, "Equipos")
ws.cell(r, 2, df.equipo.nunique())
r += 1
ws.cell(r, 1, "Atributos por jugador")
ws.cell(r, 2, len(ATTRS))
r += 2
ws.cell(r, 1, "Posiciones por fuente").font = Font(bold=True)
r += 1
for j, h in enumerate(["Fuente", "Liga", "Perfiles"], 1):
    c = ws.cell(r, j, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r += 1
for (src, liga), n in df.groupby(["src", "liga"]).size().items():
    vals = [src, liga, int(n)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.border = border
    r += 1
r += 1
ws.cell(r, 1, "Subposiciones por liga").font = Font(bold=True)
r += 1
for j, h in enumerate(["Liga", "Subposición", "Perfiles"], 1):
    c = ws.cell(r, j, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r += 1
for (liga, sub), n in df.groupby(["liga", "subposicion"]).size().items():
    vals = [liga, sub, int(n)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.border = border
    r += 1
r += 1
ws.cell(r, 1, "Notas del modelo").font = Font(bold=True, size=12, color="123B5D")
r += 1
notes = [
    "Los 37 atributos son estimaciones generadas a partir de posición, goles, edad, nivel del equipo, liga y aleatoriedad controlada. No son datos oficiales de rendimiento.",
    'La posición base está documentada para 175 perfiles (17,9%); 801 (82,1%) son estimaciones. La subposición Central/Lateral se estima cuando la fuente solo indica "Defensa".',
    "El overall se calcula con pesos específicos por demarcación (Portero, Central, Lateral, Medio, Delantero). Los goles solo refuerzan definición, flick, PC y stroke; no convierten a un defensor en delantero.",
    "Los especialistas de set pieces se asignan por el máximo atributo del equipo: inyector, stopper, drag flicker, desviador, primer runner, portero de PC, tirador de stroke, portero de stroke, aéreo y faltas.",
    "El editor del juego permite corregir posición y atributos; los cambios se guardan en localStorage del navegador y no modifican este Excel.",
    'Las posiciones oficiales incluyen fuente, URL y fecha en la hoja "Fuentes posiciones".',
    "Para una auditoría completa se recomienda validar los 175 perfiles oficiales y priorizar la revisión manual de los jugadores con más goles y de los centrales con perfil de drag flicker.",
]
for n in notes:
    ws.cell(r, 1, "• " + n).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 42
    r += 1
ws.column_dimensions["A"].width = 55
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 14

# Refresh Cobertura posiciones
ws = wb["Cobertura posiciones"]
ws.delete_rows(1, ws.max_row)
header_fill = PatternFill("solid", fgColor="123B5D")
header_font = Font(color="FFFFFF", bold=True, size=10)
r = 1
for j, h in enumerate(["Fuente", "Liga", "Perfiles"], 1):
    c = ws.cell(r, j, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r += 1
for (src, liga), n in df.groupby(["src", "liga"]).size().items():
    for j, v in enumerate([src, liga, int(n)], 1):
        c = ws.cell(r, j, v)
        c.border = border
    r += 1
for j, v in enumerate(["TOTAL", "", len(df)], 1):
    c = ws.cell(r, j, v)
    c.border = border
    c.font = Font(bold=True)
r += 2
for j, h in enumerate(["Posición", "Perfiles"], 1):
    c = ws.cell(r, j, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r += 1
for pos, n in df["pos"].value_counts().items():
    for j, v in enumerate([pos, int(n)], 1):
        c = ws.cell(r, j, v)
        c.border = border
    r += 1
r += 1
for j, h in enumerate(["Confianza", "Perfiles"], 1):
    c = ws.cell(r, j, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r += 1
for conf, n in df["conf"].value_counts().items():
    for j, v in enumerate([conf, int(n)], 1):
        c = ws.cell(r, j, v)
        c.border = border
    r += 1
ws.cell(
    r + 1, 1, 'Actualizado 2026-10-19. El detalle por jugador y URL está en la hoja "Fuentes posiciones".'
).font = Font(italic=True, color="555555")
ws.column_dimensions["A"].width = 45
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 12

# Sheet Fuentes posiciones
ws = wb.create_sheet("Fuentes posiciones")
headers = [
    "Liga",
    "Equipo",
    "Jugador",
    "Perfil",
    "Posición",
    "Subposición",
    "Fuente",
    "Confianza",
    "URL",
    "Fecha",
    "Nota",
]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
r = 2
for _, d in df.iterrows():
    if str(d.get("src", "")) == "Estimada" and not str(d.get("pos_url", "")).strip():
        continue
    vals = [
        d["liga"],
        d["equipo"],
        d["jugador"],
        d["perfil"],
        d["pos"],
        d["subposicion"],
        d["src"],
        d["conf"],
        d.get("pos_url", ""),
        d.get("pos_fecha", ""),
        d.get("pos_nota", ""),
    ]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = border
        c.alignment = left if j in (3, 7, 9, 11) else center
    r += 1
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:K{r - 1}"
for j, w in enumerate([12, 30, 32, 10, 12, 12, 28, 10, 55, 12, 35], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# Sheet Fuentes atributos (append)
if "Fuentes atributos" in wb.sheetnames:
    del wb["Fuentes atributos"]
ws = wb.create_sheet("Fuentes atributos")
headers = ["Fuente", "Uso", "URL"]
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = border
rows = [
    ("RFEH / iSquad", "Plantillas, goles, edades y actas de porteros.", "https://www.rfeh.es/"),
    ("RFEH Circular 1 2026-27", "Formato DHA/DHB, ascensos, descensos y playout.", "https://www.rfeh.es/"),
    ("Reglamento FIH", "Reglas de penalty corner, stroke, 3D, aerial y saques.", "https://www.fih.hockey/"),
    ("Páginas oficiales de clubes", "Posiciones oficiales de jugadores (175 perfiles).", "Varias"),
    ("Diari de Terrassa", "Posiciones oficiales de CD Terrassa.", "https://www.diarideterrassa.com/"),
    ("Modelo propio", "Atributos, subposiciones estimadas, arquetipos y especialistas.", "Este proyecto"),
]
for i, row in enumerate(rows, 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.border = border
        c.alignment = left if j != 1 else center
for j, w in enumerate([28, 70, 45], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

wb.save(OUT)
print("saved", OUT)
print("csv", CSV_OUT)
print("sheets", wb.sheetnames)
